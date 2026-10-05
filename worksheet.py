from openpyxl import Workbook

questions = [
    "Apply regularization to illustrate how it helps reduce overfitting in a Multi-Layer Perceptron.",
    "Analyze how the choice of an activation function affects the training of a deep neural network.",
    "Differentiate between CNN and LSTM based on the type of data and learning tasks for which they are suitable.",
    "Apply the Q-learning update mechanism to determine the improved action-value estimate using the reward and next-state information.",
    "Analyze the role of experience replay and target networks in improving the stability of a DQN.",
    "Analyze the importance of model serialization using Pickle when deploying a trained scikit-learn model in a web application.",
    "Apply the Flask deployment workflow to identify the main steps for serving predictions from a trained machine-learning model.",
    "Apply the backpropagation algorithm to show how the weights of a Multi-Layer Perceptron are updated.",
    "Differentiate between Stochastic Gradient Descent and Batch Gradient Descent based on their weight-update process.",
    "Analyze how the choice of loss function differs for classification and regression tasks.",
    "Compare a Dueling Deep Q-Network with a conventional Deep Q-Network with respect to action-value estimation.",
    "Apply the concept of multi-step learning to illustrate how future rewards are included in a reinforcement learning update.",
    "Analyze the main steps required to convert a trained regression model into a Flask-based web application.",
    "Evaluate AWS and Google Cloud as platforms for deploying a machine-learning web application.",
    "Analyze the working of an LSTM network for sequential data. Describe how its gating mechanism helps overcome the limitations of a conventional RNN, with a suitable architecture diagram.",
    "Analyze the Q-learning process for a suitable reinforcement learning problem. Explain the interaction between the agent, environment, state, action and reward, and show how Q-values are updated to select an appropriate action.",
    "Analyze the Q-learning process for a suitable reinforcement learning problem. Explain the interaction between the agent, environment, state, action and reward, and show how Q-values are updated to select an appropriate action.",
    "Design a workflow for deploying a trained scikit-learn classification model as a Flask web application. Include model serialization using Pickle, prediction handling, SQLite data storage and deployment on a public or cloud server.",
    "Analyze the complete backpropagation process in a Multi-Layer Perceptron. Describe forward propagation, loss calculation, gradient computation and weight updating with a suitable architecture diagram.",
    "Analyze how an Autoencoder performs feature learning and dimensionality reduction. Describe the roles of the encoder, latent representation and decoder with a suitable architecture diagram.",
    "Evaluate the advantages of a Dueling Deep Q-Network over a conventional Deep Q-Network. Explain how separate estimation of state value and action advantage can improve action-value estimation. Draw a suitable architecture diagram.",
    "Design a simple workflow for deploying a trained regression model as a Flask-based web application. Include model loading, user input, prediction generation and deployment on AWS or Google Cloud."
]

wb = Workbook()
ws = wb.active
ws.title = "Questions"

row = 1

for i, question in enumerate(questions, start=1):
    ws.cell(row=row, column=1, value=i)
    ws.cell(row=row, column=2, value=question)

    row += 2   # leaves one blank row after every question

ws.column_dimensions["A"].width = 8
ws.column_dimensions["B"].width = 80

wb.save("questions.xlsx")

print("questions.xlsx created successfully!")
